import io
import re
from typing import Optional
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, joinedload
import openpyxl

from app.database import get_db
from app import models, schemas, auth, crud
from app.utils.excel_generator import generate_books_excel

router = APIRouter(prefix="/api/reports", tags=["Reports"])


def _collect_rows(db: Session, series_id: Optional[int], series_code: Optional[str],
                   sub_series_id: Optional[int],
                   author: Optional[str], search: Optional[str],
                   order_by: str = "series") -> tuple[list[dict], str]:
    q = db.query(models.Book).options(
        joinedload(models.Book.series), joinedload(models.Book.sub_series), joinedload(models.Book.copies)
    )
    report_title = "Library Report - All Series"

    if series_id:
        q = q.filter(models.Book.series_id == series_id)
        series = db.query(models.Series).filter(models.Series.id == series_id).first()
        if series:
            report_title = f"Library Report - {series.code} ({series.name})"
    if series_code:
        q = q.join(models.Series).filter(models.Series.code == series_code.upper())
        report_title = f"Library Report - {series_code.upper()}"
    if sub_series_id:
        q = q.filter(models.Book.sub_series_id == sub_series_id)
        sub = db.query(models.SubSeries).filter(models.SubSeries.id == sub_series_id).first()
        if sub:
            report_title += f" / {sub.name}"
    if author:
        q = q.filter(models.Book.author.ilike(f"%{author}%"))
    if search:
        like = f"%{search}%"
        q = q.filter((models.Book.title.ilike(like)) | (models.Book.author.ilike(like)))

    books = q.all()

    # One row per physical copy. In Series Order, each book's copies stay
    # grouped under its serial number (A-12, A-12(2), A-13, ...). In Latest
    # Added Order, rows are sorted purely by when each copy was added,
    # newest first - so a freshly added copy of an older book (e.g. a new
    # A-12(2)) rises to the top instead of being buried right after A-12,
    # without ever changing the serial numbering scheme.
    #
    # IMPORTANT: a book can legitimately have zero copies left (e.g. its
    # only copy was deleted via "Delete Copy" - marked lost/withdrawn and
    # removed - without deleting the book's catalog record). Looping only
    # over b.copies would silently drop that book from the report entirely,
    # even though it still exists in the catalog and shows up on the Books
    # page. That's the root cause of exports appearing to "lose" books as a
    # series accumulates copy-level edits over time. Every book that matches
    # the filters MUST contribute at least one row.
    rows = []
    for b in books:
        if b.copies:
            for c in b.copies:
                rows.append({
                    "serial": crud.display_serial_for_book(b, c.copy_number),
                    "title": b.title,
                    "author": b.author,
                    "category": b.series.name,
                    "sub_category": b.sub_series.name if b.sub_series else "—",
                    "_series_key": (b.series_id, b.base_serial, c.copy_number),
                    "_latest_key": c.created_at or b.created_at,
                })
        else:
            rows.append({
                "serial": crud.display_serial_for_book(b),
                "title": b.title,
                "author": b.author,
                "category": b.series.name,
                "sub_category": b.sub_series.name if b.sub_series else "—",
                "_series_key": (b.series_id, b.base_serial, 0),
                "_latest_key": b.created_at,
            })

    if order_by == "latest":
        rows.sort(key=lambda r: r["_latest_key"], reverse=True)
    else:
        rows.sort(key=lambda r: r["_series_key"])

    for r in rows:
        r.pop("_series_key", None)
        r.pop("_latest_key", None)

    return rows, report_title


ORDER_LABELS = {"series": "Series Order", "latest": "Latest Added Order"}


def _content_disposition(filename: str) -> str:
    """
    Build a Content-Disposition header that survives non-ASCII (e.g. Kannada)
    report titles. HTTP header values must be Latin-1 encodable, so a filename
    containing Kannada characters raised a UnicodeEncodeError and crashed the
    download (both PDF and Excel) before any bytes were sent. We now send an
    ASCII-safe fallback name plus the proper RFC 5987/6266 filename* parameter
    with the real Unicode name, so browsers save the file with the correct
    (Kannada) name while older clients still get a sane ASCII fallback.
    """
    ascii_fallback = re.sub(r"[^A-Za-z0-9._-]+", "_", filename).strip("_") or "report"
    return f"attachment; filename=\"{ascii_fallback}\"; filename*=UTF-8''{quote(filename)}"


@router.get("/books/excel")
def export_books_excel(
    series_id: Optional[int] = None, series_code: Optional[str] = None,
    sub_series_id: Optional[int] = None,
    author: Optional[str] = None, search: Optional[str] = None,
    order_by: str = Query(
        "series", pattern="^(series|latest)$",
        description="'series' groups by serial number; 'latest' shows newest-added copies first.",
    ),
    db: Session = Depends(get_db), _user: models.User = Depends(auth.get_current_user),
):
    rows, title = _collect_rows(db, series_id, series_code, sub_series_id, author, search, order_by)
    xlsx_bytes = generate_books_excel(rows, title, order_label=ORDER_LABELS[order_by])
    filename = title.replace(" ", "_").replace("(", "").replace(")", "") + ".xlsx"
    return StreamingResponse(
        io.BytesIO(xlsx_bytes),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": _content_disposition(filename)},
    )


@router.post("/import/excel")
async def import_books_excel(
    series_id: int = Query(..., description="Series to import these rows into"),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _admin: models.User = Depends(auth.require_admin),
):
    """
    Bulk import books from an Excel file into the given series.
    Required columns (header row, any order, case-insensitive): Title, Author, Category.
    Optional columns: Language, Publisher, Year, ISBN, Notes.
    Column names must match exactly (aside from case/leading-trailing spaces) -
    if a required column isn't found, the import is rejected up front with a
    message naming exactly which column header is missing, so it can be
    renamed in the file and re-uploaded, rather than guessing and silently
    importing into the wrong field.
    Duplicate (title, author) pairs are automatically added as additional copies.
    """
    series = db.query(models.Series).filter(models.Series.id == series_id).first()
    if not series:
        raise HTTPException(status_code=404, detail="Series not found")

    content = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read Excel file. Please upload a valid .xlsx file.")

    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    raw_header = list(next(rows_iter))

    def _norm(h) -> str:
        return str(h).strip().lower() if h else ""

    norm_header = [_norm(h) for h in raw_header]

    def col(name: str) -> Optional[int]:
        return norm_header.index(name) if name in norm_header else None

    REQUIRED_COLUMNS = ["title", "author", "category"]

    missing = [name.capitalize() for name in REQUIRED_COLUMNS if col(name) is None]
    if missing:
        found = ", ".join(str(h).strip() for h in raw_header if h and str(h).strip()) or "(none)"
        raise HTTPException(
            status_code=400,
            detail=(
                f"Column name(s) not recognized: {', '.join(missing)}. "
                "The Excel file's header row must contain columns named exactly "
                "'Title', 'Author', and 'Category' (case-insensitive). "
                f"Columns found in your file: {found}. "
                "Please rename the column header(s) to match and re-upload."
            ),
        )

    idx_title, idx_author, idx_subseries = col("title"), col("author"), col("category")
    idx_lang, idx_pub, idx_year, idx_isbn, idx_notes = (
        col("language"), col("publisher"), col("year"), col("isbn"), col("notes"))

    created, copies_added, errors = 0, 0, []
    for row_num, row in enumerate(rows_iter, start=2):
        if not row or not row[idx_title]:
            continue
        try:
            payload = schemas.BookCreate(
                series_id=series_id,
                sub_series_name=str(row[idx_subseries]).strip() if idx_subseries is not None and row[idx_subseries] else None,
                title=str(row[idx_title]).strip(),
                author=str(row[idx_author]).strip() if idx_author is not None and row[idx_author] else "Unknown",
                language=str(row[idx_lang]).strip() if idx_lang is not None and row[idx_lang] else None,
                publisher=str(row[idx_pub]).strip() if idx_pub is not None and row[idx_pub] else None,
                year_published=int(row[idx_year]) if idx_year is not None and row[idx_year] else None,
                isbn=str(row[idx_isbn]).strip() if idx_isbn is not None and row[idx_isbn] else None,
                notes=str(row[idx_notes]).strip() if idx_notes is not None and row[idx_notes] else None,
            )
            _, was_new = crud.create_book_or_add_copy(db, payload)
            created += 1 if was_new else 0
            copies_added += 0 if was_new else 1
        except Exception as e:
            errors.append(f"Row {row_num}: {str(e)}")

    return {
        "new_books_created": created,
        "additional_copies_added": copies_added,
        "errors": errors,
    }


@router.post("/import/taranga-excel")
async def import_taranga_excel(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _admin: models.User = Depends(auth.require_admin),
):
    """
    Bulk import Taranga entries from an Excel file.

    Unlike the general book importer, Taranga only ever needs Title and
    Month - there's no Author/Category/etc to fill in, and no series to
    choose (every row is auto-assigned to the fixed Taranga series, same as
    the "Add Taranga" quick-entry form). Required column (header row, any
    order, case-insensitive): Title. Optional column: Month. Any other
    columns present (e.g. a leftover Author/Category column from reusing a
    book template) are simply ignored rather than rejected, so a librarian
    can reuse a familiar spreadsheet layout without stripping columns first.
    Every valid row becomes its own new Taranga entry (no dedup/merge,
    matching the single-entry "Add Taranga" behavior) with its own
    auto-assigned serial number.
    """
    content = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read Excel file. Please upload a valid .xlsx file.")

    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        raw_header = list(next(rows_iter))
    except StopIteration:
        raise HTTPException(status_code=400, detail="The Excel file is empty - no header row was found.")

    def _norm(h) -> str:
        return str(h).strip().lower() if h else ""

    norm_header = [_norm(h) for h in raw_header]

    def col(name: str) -> Optional[int]:
        return norm_header.index(name) if name in norm_header else None

    idx_title = col("title")
    if idx_title is None:
        found = ", ".join(str(h).strip() for h in raw_header if h and str(h).strip()) or "(none)"
        raise HTTPException(
            status_code=400,
            detail=(
                "Column name not recognized: Title. The Excel file's header row must contain "
                "a column named exactly 'Title' (case-insensitive). "
                f"Columns found in your file: {found}. "
                "Please rename the column header and re-upload."
            ),
        )
    idx_month = col("month")

    created, errors = 0, []
    for row_num, row in enumerate(rows_iter, start=2):
        if not row or all(v is None or str(v).strip() == "" for v in row):
            continue  # skip fully blank rows
        title_val = row[idx_title] if idx_title < len(row) else None
        if not title_val or not str(title_val).strip():
            errors.append(f"Row {row_num}: Title is required and was left blank.")
            continue
        try:
            payload = schemas.TarangaCreate(
                title=str(title_val).strip(),
                month=str(row[idx_month]).strip() if idx_month is not None and idx_month < len(row) and row[idx_month] else None,
            )
            crud.create_taranga(db, payload)
            created += 1
        except Exception as e:
            errors.append(f"Row {row_num}: {str(e)}")

    return {
        "new_taranga_created": created,
        "errors": errors,
    }
